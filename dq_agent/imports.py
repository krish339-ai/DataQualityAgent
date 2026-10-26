"""Bounded, local-only table importers. No model, network, or generated code.

The engine validates a string table. Its versioned manifest retains source types
and positions so that this text view is never mistaken for a lossless workbook
or typed-JSON export. Row indices in type metadata are zero-based table rows.
"""
from __future__ import annotations

import csv
import io
import json
import posixpath
import re
import time
import zipfile
import zlib
from dataclasses import dataclass
from datetime import date, datetime, time as daytime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import PurePosixPath
from typing import Any

import pandas as pd
from defusedxml import ElementTree as SafeXML
from defusedxml.common import DefusedXmlException
from openpyxl import load_workbook
from openpyxl.styles.numbers import is_datetime
from openpyxl.utils.cell import column_index_from_string, get_column_letter


MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 200
MAX_CELLS = 1_000_000
MAX_DECODED_BYTES = 50 * 1024 * 1024
MAX_ZIP_MEMBERS = 1_000
MAX_XML_DEPTH = 64
MAX_PARSE_SECONDS = 30
PREVIEW_ROWS = 100
FORMATS = {".csv": "csv", ".tsv": "tsv", ".xlsx": "xlsx", ".json": "json", ".jsonl": "jsonl", ".ndjson": "ndjson"}
_CELL_REF = re.compile(r"^([A-Z]{1,3})([1-9][0-9]*)$")
_NUMBER = re.compile(r"-?(?:0|[1-9][0-9]*)(?:\.[0-9]+)?(?:[eE][+-]?[0-9]+)?\Z")


class ImportValidationError(ValueError):
    """A structural or resource issue that must be resolved before creating a job."""


@dataclass(frozen=True)
class ParsedImport:
    csv_bytes: bytes
    manifest: dict[str, Any]
    preview: pd.DataFrame


class _Budget:
    def __init__(self):
        self.started = time.monotonic()
        self.decoded = 0

    def tick(self):
        if time.monotonic() - self.started > MAX_PARSE_SECONDS:
            raise ImportValidationError("Import inspection exceeded its time budget. Split the file into smaller tables.")

    def text(self, value: str):
        self.tick()
        try:
            size = len(value.encode("utf-8"))
        except UnicodeEncodeError as exc:
            raise ImportValidationError("Text contains an invalid Unicode character. Save valid UTF-8 text.") from exc
        if "\x00" in value:
            raise ImportValidationError("The file contains a NUL character. Supply ordinary text values.")
        self.decoded += size
        if self.decoded > MAX_DECODED_BYTES:
            raise ImportValidationError("Decoded cell values exceed the 50 MiB limit. Split the table.")


def _identity(data: bytes, filename: str) -> tuple[str, str]:
    if not isinstance(data, bytes) or not data:
        raise ImportValidationError("Choose a nonempty file.")
    if len(data) > MAX_BYTES:
        raise ImportValidationError("Files must be no larger than 10 MiB.")
    if not isinstance(filename, str) or not filename.strip():
        raise ImportValidationError("Provide a filename with a supported extension.")
    # A display name only; it is never used as a storage path.
    name = filename.replace("\\", "/").rsplit("/", 1)[-1]
    if len(name) > 255 or any(ord(c) < 32 for c in name):
        raise ImportValidationError("Use a filename of at most 255 characters without control characters.")
    try:
        name.encode("utf-8")
    except UnicodeEncodeError as exc:
        raise ImportValidationError("The filename contains an invalid Unicode character. Rename the file.") from exc
    fmt = FORMATS.get(PurePosixPath(name).suffix.lower())
    if fmt is None:
        raise ImportValidationError("Supported files are CSV, TSV, XLSX, flat JSON, JSONL, and NDJSON. Save a supported copy.")
    return name, fmt


def _decode(data: bytes) -> str:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ImportValidationError("Text files must use UTF-8 encoding; a UTF-8 BOM is supported.") from exc
    if "\x00" in text:
        raise ImportValidationError("Text contains a NUL character. Save a plain UTF-8 file.")
    if len(text.encode("utf-8")) > MAX_DECODED_BYTES:
        raise ImportValidationError("Decoded text exceeds the 50 MiB limit.")
    return text


def _shape(row_count: int, column_count: int):
    if column_count > MAX_COLUMNS:
        raise ImportValidationError("A table may have at most 200 columns.")
    if row_count > MAX_ROWS:
        raise ImportValidationError("A table may have at most 100,000 data rows.")
    if row_count * column_count > MAX_CELLS:
        raise ImportValidationError("A table may contain at most 1,000,000 data cells. Split the table.")


def _headers(columns: list[str], budget: _Budget):
    _shape(0, len(columns))
    if not columns or any(not isinstance(c, str) or not c.strip() for c in columns):
        raise ImportValidationError("The table needs nonblank text column headers in its first row.")
    if len(columns) != len(set(columns)):
        raise ImportValidationError("Column headers must be unique. Rename duplicate columns before importing.")
    for column in columns:
        budget.text(column)


def _typed(types: dict, column: str, kind: str, row: int):
    if kind != "string":
        types.setdefault(column, {}).setdefault(kind, []).append(row)


def _finish(name, fmt, columns, rows, source_locations, *, options, budget,
            types=None, numeric_lexemes=None, warnings=None, sheet=None) -> ParsedImport:
    if not rows:
        raise ImportValidationError("The table must contain at least one data row.")
    _shape(len(rows), len(columns))
    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    writer.writerow(columns)
    writer.writerows(rows)
    normalized = output.getvalue().encode("utf-8")
    # The existing engine has the same bounded internal CSV contract.
    if len(normalized) > MAX_BYTES:
        raise ImportValidationError("The normalized table exceeds the 10 MiB working-table limit. Split the file.")
    budget.tick()
    manifest = {
        "version": 1, "format": fmt, "original_filename": name,
        "options": options, "sheet": sheet,
        "row_count": len(rows), "column_count": len(columns), "columns": columns,
        "types": {"default": "string", "by_column": types or {}},
        "numeric_lexemes": numeric_lexemes or {}, "source_locations": source_locations,
        "warnings": warnings or [],
    }
    return ParsedImport(normalized, manifest, pd.DataFrame(rows[:PREVIEW_ROWS], columns=columns, dtype=str))


def _delimited(data, name, fmt, budget):
    delimiter = "\t" if fmt == "tsv" else ","
    # The standard CSV parser's otherwise small default rejects legitimate fields.
    previous_field_limit = csv.field_size_limit(MAX_BYTES)
    reader = csv.reader(io.StringIO(_decode(data), newline=""), delimiter=delimiter, strict=True)
    try:
        columns = next(reader, None)
        _headers(columns or [], budget)
        rows, locations, blank_count = [], [], 0
        previous_line = reader.line_num
        for row in reader:
            budget.tick()
            _shape(len(rows) + 1, len(columns))
            start, end = previous_line + 1, reader.line_num
            previous_line = end
            if row == []:
                row = [""] * len(columns)
                blank_count += 1
            if len(row) != len(columns):
                raise ImportValidationError(f"Data row {len(rows) + 1} has {len(row)} fields; expected {len(columns)}. Correct the delimiter or row structure.")
            for value in row:
                budget.text(value)
            rows.append(row)
            locations.append([start, end])
    except csv.Error as exc:
        raise ImportValidationError("Delimited text has malformed quoting or an oversized field. Correct it before importing.") from exc
    finally:
        csv.field_size_limit(previous_field_limit)
    warnings = [f"Retained {blank_count} empty physical record(s) as blank data rows."] if blank_count else []
    return _finish(name, fmt, columns, rows, {"kind": "delimited", "rows": locations, "columns": columns},
                   options={"encoding": "utf-8-sig", "delimiter": delimiter, "header_row": 1},
                   budget=budget, warnings=warnings)


class _JsonNumber(str):
    """Keep the source token instead of converting through binary float."""


def _number(token: str) -> str:
    try:
        if not _NUMBER.fullmatch(token) or not Decimal(token).is_finite():
            raise InvalidOperation
    except (InvalidOperation, ValueError) as exc:
        raise ImportValidationError("Numbers must be finite JSON-style numeric values.") from exc
    return token


def _pairs(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ImportValidationError("JSON contains a duplicate object key. Use unique keys in every record.")
        result[key] = value
    return result


def _constant(_):
    raise ImportValidationError("JSON NaN and Infinity are unsupported. Supply finite numbers or explicit nulls.")


def _load_json(text):
    try:
        return json.loads(text, object_pairs_hook=_pairs, parse_int=_JsonNumber,
                          parse_float=_JsonNumber, parse_constant=_constant)
    except (json.JSONDecodeError, RecursionError) as exc:
        raise ImportValidationError("JSON is malformed or too deeply nested. Use flat record objects.") from exc


def _json_table(data, name, fmt, budget):
    text = _decode(data)
    if fmt == "json":
        records = _load_json(text)
        if not isinstance(records, list):
            raise ImportValidationError("A JSON file must contain an array of flat record objects. Use JSONL for one object per line.")
        _shape(len(records), 0)
        locations = list(range(len(records)))
    else:
        lines = text.split("\n")
        if lines and lines[-1] == "":
            lines.pop()  # A final line terminator is not an extra record.
        records, locations = [], []
        for line_number, line in enumerate(lines, 1):
            budget.tick()
            if not line.strip():
                raise ImportValidationError(f"JSONL/NDJSON line {line_number} is blank. Remove it or use an explicit object; blank lines are not silently skipped.")
            _shape(len(records) + 1, 0)
            records.append(_load_json(line))
            locations.append(line_number)
    columns, seen = [], set()
    for record in records:
        budget.tick()
        if not isinstance(record, dict):
            raise ImportValidationError("Each JSON record must be a flat object with scalar values.")
        for column, value in record.items():
            if isinstance(value, (dict, list)):
                raise ImportValidationError("Nested JSON objects and arrays are unsupported. Supply a flat table with one scalar value per field.")
            if column not in seen:
                columns.append(column)
                seen.add(column)
                _shape(len(records), len(columns))
    _headers(columns, budget)
    rows, types, numbers = [], {}, {}
    for index, record in enumerate(records):
        row = []
        for column in columns:
            if column not in record:
                value, kind = "", "missing"
            elif record[column] is None:
                value, kind = "", "null"
            elif isinstance(record[column], _JsonNumber):
                value, kind = _number(record[column]), "number"
                numbers.setdefault(column, {})[str(index)] = str(value)
            elif isinstance(record[column], bool):
                value, kind = ("true" if record[column] else "false"), "boolean"
            else:
                value, kind = record[column], "string"
            budget.text(value)
            row.append(str(value))
            _typed(types, column, kind, index)
        rows.append(row)
    warnings = ["Validation uses a text view. Original JSON scalar types and exact numeric tokens are retained in the import manifest; CSV exports are text tables."]
    missing = sum(len(kinds.get("missing", [])) for kinds in types.values())
    nulls = sum(len(kinds.get("null", [])) for kinds in types.values())
    if missing or nulls:
        warnings.append(f"The text view shows {missing} missing field(s) and {nulls} null value(s) as blank; the manifest distinguishes them from empty strings.")
    return _finish(name, fmt, columns, rows,
                   {"kind": "json_array" if fmt == "json" else "json_lines", "rows": locations, "columns": columns},
                   options={"encoding": "utf-8-sig", "record_structure": "array" if fmt == "json" else "one-object-per-line",
                            "column_order": "first-seen", "null_text": "", "missing_text": "", "numbers": "exact-source-token"},
                   budget=budget, types=types, numeric_lexemes=numbers, warnings=warnings)


def _local(tag):
    return tag.rsplit("}", 1)[-1]


def _xml(payload: bytes, budget: _Budget):
    root, depth = None, 0
    try:
        for event, element in SafeXML.iterparse(io.BytesIO(payload), events=("start", "end"),
                                               forbid_dtd=True, forbid_entities=True, forbid_external=True):
            budget.tick()
            if event == "start":
                depth += 1
                if root is None:
                    root = element
                if depth > MAX_XML_DEPTH:
                    raise ImportValidationError("Workbook XML is too deeply nested. Re-save a simple values-only workbook.")
            else:
                depth -= 1
    except (SafeXML.ParseError, DefusedXmlException) as exc:
        raise ImportValidationError("Workbook XML is malformed or contains unsafe entity/DTD declarations. Re-save a values-only XLSX.") from exc
    return root


def _archive(data: bytes, budget: _Budget) -> dict:
    """Inspect every ZIP member and XML part before openpyxl receives the file."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_MEMBERS:
                raise ImportValidationError("Workbook has too many ZIP parts. Save a smaller values-only workbook.")
            names, expanded, workbook, relationships = set(), 0, None, None
            for entry in entries:
                budget.tick()
                part = entry.filename
                if part in names or "\\" in part or part.startswith("/") or ".." in PurePosixPath(part).parts:
                    raise ImportValidationError("Workbook ZIP has duplicate or unsafe member paths. Re-save the workbook.")
                names.add(part)
                if entry.flag_bits & 1:
                    raise ImportValidationError("Encrypted workbooks are unsupported. Save an unencrypted values-only XLSX.")
                if entry.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}:
                    raise ImportValidationError("Workbook uses unsupported compression. Save a standard XLSX copy.")
                expanded += entry.file_size
                if expanded > MAX_DECODED_BYTES:
                    raise ImportValidationError("Workbook ZIP expansion exceeds 50 MiB. Split or simplify the workbook.")
                lower = part.lower()
                if any(marker in lower for marker in ("vbaproject", "externallinks/", "connections.xml", "querytables/", "macrosheets/", "xl/activex/", "xl/embeddings/")):
                    raise ImportValidationError("Macros, external links, and refresh connections are unsupported. Save a values-only XLSX without them.")
                # Bounded reads verify CRC, declared length, and compression validity.
                with archive.open(entry) as handle:
                    payload = handle.read(MAX_DECODED_BYTES + 1)
                if len(payload) != entry.file_size or len(payload) > MAX_DECODED_BYTES:
                    raise ImportValidationError("Workbook ZIP member size is invalid or exceeds the limit.")
                if lower.endswith((".xml", ".rels")):
                    root = _xml(payload, budget)
                    for element in root.iter():
                        tag = _local(element.tag)
                        if tag in {"f", "formula", "formula1", "formula2"}:
                            raise ImportValidationError("Workbook contains formulas. Copy the table and paste values only into a new XLSX before importing.")
                        if tag == "Relationship" and element.attrib.get("TargetMode", "").lower() == "external":
                            raise ImportValidationError("Workbook contains external links. Remove links/connections and save a values-only XLSX.")
                        if tag in {"Override", "Default"} and any(s in element.attrib.get("ContentType", "").lower() for s in ("macroenabled", "vbaproject")):
                            raise ImportValidationError("Macro-enabled workbooks are unsupported. Save a values-only XLSX.")
                    if part == "xl/workbook.xml":
                        workbook = root
                    elif part == "xl/_rels/workbook.xml.rels":
                        relationships = root
            if workbook is None or relationships is None or "[Content_Types].xml" not in names:
                raise ImportValidationError("This file is not a complete XLSX workbook. Save a new .xlsx copy.")
            rels = {r.attrib.get("Id"): r.attrib for r in relationships if _local(r.tag) == "Relationship"}
            sheets = []
            for element in workbook.iter():
                if _local(element.tag) != "sheet":
                    continue
                relation_id = next((v for k, v in element.attrib.items() if _local(k) == "id"), None)
                relation = rels.get(relation_id, {})
                if not relation.get("Type", "").endswith("/worksheet"):
                    raise ImportValidationError("Only standard worksheets are supported. Save the data table in a separate XLSX.")
                target = relation.get("Target", "")
                path = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join("xl", target))
                if not path.startswith("xl/worksheets/") or path not in names:
                    raise ImportValidationError("Workbook worksheet references are invalid. Re-save the workbook.")
                sheets.append({"name": element.attrib.get("name", ""), "state": element.attrib.get("state", "visible"), "path": path})
            if not sheets or any(not s["name"] for s in sheets) or len({s["name"] for s in sheets}) != len(sheets):
                raise ImportValidationError("Workbook needs uniquely named worksheets.")
            return {"sheets": sheets, "expanded_bytes": expanded}
    except ImportValidationError:
        raise
    except (zipfile.BadZipFile, RuntimeError, OSError, EOFError, NotImplementedError, ValueError, zlib.error) as exc:
        raise ImportValidationError("XLSX is corrupt, encrypted, or uses unsupported compression. Save a new values-only XLSX.") from exc


def inspect_upload(data: bytes, filename: str) -> dict:
    """Identify a supported file and enumerate sheets without creating any job."""
    name, fmt = _identity(data, filename)
    result = {"format": fmt, "original_filename": name, "sheets": [], "requires_sheet": False}
    if fmt == "xlsx":
        info = _archive(data, _Budget())
        result.update(sheets=[s["name"] for s in info["sheets"]],
                      sheet_states={s["name"]: s["state"] for s in info["sheets"]},
                      requires_sheet=len(info["sheets"]) > 1)
    else:
        _decode(data)
    return result


def _zero_padding(number_format: str) -> bool:
    for part in number_format.split(";"):
        clean = re.sub(r'"[^"]*"|\\.|\[[^\]]*\]|_.|\*.', "", part)
        integer = re.split(r"[.Ee]", clean, maxsplit=1)[0]
        if integer.count("0") > 1:
            return True
    return False


def _date_kind(number_format: str) -> str | None:
    # Excel format tokens are case-insensitive.
    return is_datetime(number_format.lower())


def _native_date(token: str):
    """Validate the complete ISO token, including its explicit timezone."""
    try:
        if "T" in token or " " in token:
            return datetime.fromisoformat(token)
        if ":" in token:
            return daytime.fromisoformat(token)
        return date.fromisoformat(token)
    except ValueError as exc:
        raise ImportValidationError("A native Excel date contains an invalid ISO value. Correct the date or store it as text.") from exc


def _sheet_structure(root, budget):
    numeric, native_dates, occupied, hidden_rows = {}, {}, set(), 0
    max_row, max_column, hidden_columns = 0, 0, False
    previous_row, previous_column = 0, 0
    for element in root.iter():
        budget.tick()
        tag = _local(element.tag)
        if tag == "mergeCell":
            raise ImportValidationError("The selected worksheet contains merged cells. Unmerge them and supply explicit values in a rectangular table.")
        if tag == "row":
            try:
                row = int(element.attrib["r"])
            except (KeyError, ValueError) as exc:
                raise ImportValidationError("Worksheet row positions are invalid. Re-save the workbook.") from exc
            if row <= previous_row:
                raise ImportValidationError("Worksheet row positions must be positive, unique, and ordered. Re-save the workbook.")
            previous_row, previous_column = row, 0
            max_row = max(max_row, row)
            _shape(max_row - 1, max_column)
            hidden_rows += element.attrib.get("hidden") in {"1", "true"}
        elif tag == "col" and element.attrib.get("hidden") in {"1", "true"}:
            hidden_columns = True
        elif tag == "c":
            coordinate = element.attrib.get("r", "")
            match = _CELL_REF.fullmatch(coordinate)
            if not match or coordinate in occupied:
                raise ImportValidationError("Worksheet cell positions are invalid or duplicated. Re-save the workbook.")
            occupied.add(coordinate)
            column, row = column_index_from_string(match[1]), int(match[2])
            if row != previous_row or column <= previous_column:
                raise ImportValidationError("Worksheet cells must match their row and be ordered. Re-save the workbook.")
            previous_column = column
            max_row, max_column = max(max_row, row), max(max_column, column)
            _shape(max_row - 1, max_column)
            cell_type = element.attrib.get("t", "n")
            raw = next((v.text for v in element if _local(v.tag) == "v"), None)
            if cell_type == "n" and raw is not None:
                numeric[coordinate] = _number(raw)
            elif cell_type == "d" and raw is not None:
                native_dates[coordinate] = _native_date(raw)
    return max_row, max_column, numeric, native_dates, hidden_rows, hidden_columns


def _xlsx(data, name, sheet, budget):
    info = _archive(data, budget)
    choices = [s["name"] for s in info["sheets"]]
    if sheet is None:
        if len(choices) != 1:
            raise ImportValidationError("This workbook has multiple worksheets. Select the worksheet to import.")
        sheet = choices[0]
    if sheet not in choices:
        raise ImportValidationError("The selected worksheet does not exist. Choose one of the listed worksheets.")
    selected = next(s for s in info["sheets"] if s["name"] == sheet)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        root = _xml(archive.read(selected["path"]), budget)
    max_row, max_column, raw_numbers, native_dates, hidden_rows, hidden_columns = _sheet_structure(root, budget)
    if max_row < 2 or max_column == 0:
        raise ImportValidationError("The selected worksheet needs a header in row 1 and at least one data row.")
    types, numbers, rows, blank_count = {}, {}, [], 0
    formatted_numeric = False
    workbook = None
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        selected_sheet = workbook[sheet]
        # Use the inspected cell extent, not a potentially false dimension hint.
        selected_sheet.reset_dimensions()
        iterator = selected_sheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_column)
        header_cells = next(iterator)
        columns = [cell.value for cell in header_cells]
        _headers(columns, budget)
        for index, cells in enumerate(iterator):
            row = []
            for position, cell in enumerate(cells):
                budget.tick()
                column = columns[position]
                coordinate = f"{get_column_letter(position + 1)}{index + 2}"
                raw = raw_numbers.get(coordinate)
                value = native_dates.get(coordinate, cell.value)
                # Excel's 1900 calendar includes the nonexistent February 29.
                # openpyxl maps that serial onto February 28; do not silently
                # make this factual date correction during an import.
                if (raw is not None and workbook.epoch.year == 1899
                        and _date_kind(cell.number_format) in {"date", "datetime"}
                        and Decimal(60) <= Decimal(raw) < Decimal(61)):
                    raise ImportValidationError(f"Cell {coordinate} uses Excel's nonexistent 1900-02-29 date. Correct the date or supply explicit text before importing.")
                if cell.data_type in {"f", "e"}:
                    raise ImportValidationError(f"Cell {coordinate} contains a formula or Excel error. Supply a values-only table with resolved values.")
                if value is None:
                    text, kind = "", "string" if cell.data_type in {"s", "inlineStr"} else "blank"
                elif isinstance(value, bool):
                    text, kind = ("true" if value else "false"), "boolean"
                elif isinstance(value, datetime):
                    # A date-only display format must not erase a stored time
                    # component, timezone, or explicitly typed ISO datetime.
                    date_only = (coordinate not in native_dates and _date_kind(cell.number_format) == "date"
                                 and value.time() == daytime(0, 0) and value.tzinfo is None)
                    kind = "date" if date_only else "datetime"
                    text = value.date().isoformat() if kind == "date" else value.isoformat()
                elif isinstance(value, (date, daytime)):
                    text, kind = value.isoformat(), "date" if isinstance(value, date) else "time"
                elif isinstance(value, timedelta):
                    raise ImportValidationError(f"Cell {coordinate} uses an Excel duration. Convert it to explicit text before importing.")
                elif isinstance(value, str):
                    text, kind = value, "string"
                elif isinstance(value, (int, float, Decimal)):
                    if _zero_padding(cell.number_format):
                        raise ImportValidationError(f"Cell {coordinate} uses a zero-padded numeric format. Store the intended identifier as text in a values-only XLSX to preserve leading zeros.")
                    if raw is None:
                        raise ImportValidationError(f"Cell {coordinate} has no recoverable numeric token. Store it as explicit text.")
                    text, kind = raw, "number"
                    formatted_numeric |= cell.number_format != "General"
                else:
                    raise ImportValidationError(f"Cell {coordinate} has an unsupported value type. Convert it to explicit text.")
                if raw is not None:
                    numbers.setdefault(column, {})[str(index)] = raw
                budget.text(text)
                row.append(text)
                _typed(types, column, kind, index)
            blank_count += all(value == "" for value in row)
            rows.append(row)
    except ImportValidationError:
        raise
    except Exception as exc:
        raise ImportValidationError("The worksheet could not be read safely. Re-save a simple values-only XLSX and try again.") from exc
    finally:
        if workbook is not None:
            workbook.close()
    warnings = ["Validation uses a text view of the selected worksheet. Dates use ISO notation; text identifiers retain leading zeros. Original types, positions, and workbook bytes are retained separately.",
                "Exports contain the reviewed data table; workbook formatting, charts, and other worksheets are not reproduced."]
    if selected["state"] != "visible":
        warnings.append("The selected worksheet is hidden in Excel; its records are included.")
    if hidden_rows or hidden_columns:
        warnings.append(f"Hidden rows/columns are included ({hidden_rows} hidden row(s)); no records were silently filtered.")
    if blank_count:
        warnings.append(f"Retained {blank_count} blank worksheet row(s) within the inspected table extent.")
    if formatted_numeric:
        warnings.append("Numeric cells use their stored numeric values; display formatting such as currency, percentage, or rounding is not applied.")
    if any(kind in {"date", "datetime", "time"} for kinds in types.values() for kind in kinds):
        warnings.append("Excel serial date/time conversion uses workbook calendar settings and millisecond precision; exact stored serial tokens remain in the manifest.")
    return _finish(name, "xlsx", columns, rows,
                   {"kind": "worksheet", "sheet": sheet, "rows": list(range(2, max_row + 1)),
                    "columns": [get_column_letter(i + 1) for i in range(max_column)]},
                   options={"header_row": 1, "range": f"A1:{get_column_letter(max_column)}{max_row}",
                            "hidden_records": "included", "numeric_values": "stored-token", "dates": "iso8601"},
                   budget=budget, types=types, numeric_lexemes=numbers, warnings=warnings, sheet=sheet)


def parse_upload(data: bytes, filename: str, *, sheet: str | None = None) -> ParsedImport:
    """Validate a complete supported table and return its bounded text view."""
    name, fmt = _identity(data, filename)
    budget = _Budget()
    if fmt == "xlsx":
        return _xlsx(data, name, sheet, budget)
    if sheet is not None:
        raise ImportValidationError("Worksheet selection applies only to XLSX files.")
    if fmt in {"csv", "tsv"}:
        return _delimited(data, name, fmt, budget)
    return _json_table(data, name, fmt, budget)

"""Verify preservation, provenance, structure rejection, and import resource bounds."""
import csv
import io
import json
import zipfile
from datetime import date, datetime, time

import pytest
from openpyxl import Workbook

import dq_agent.imports as adapters
from dq_agent.imports import ImportValidationError, inspect_upload, parse_upload


def workbook_bytes(rows=None, customize=None):
    workbook = Workbook()
    workbook.active.title = "Customers"
    for row in rows or [["id", "name"], ["001", "Ada"]]:
        workbook.active.append(row)
    if customize:
        customize(workbook)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def rewrite_zip(data, replacements=None, extras=None):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for member in source.infolist():
            value = source.read(member.filename)
            transform = (replacements or {}).get(member.filename)
            if transform:
                value = transform(value)
            target.writestr(member.filename, value)
        for name, value in (extras or {}).items():
            target.writestr(name, value)
    return output.getvalue()


def table(parsed):
    return list(csv.reader(io.StringIO(parsed.csv_bytes.decode())))


@pytest.mark.parametrize("filename,data", [
    ("customers.csv", b"id,name\n001,Ada\n002,Ben\n"),
    ("customers.tsv", b"id\tname\n001\tAda\n002\tBen\n"),
    ("customers.json", b'[{"id":"001","name":"Ada"},{"id":"002","name":"Ben"}]'),
    ("customers.jsonl", b'{"id":"001","name":"Ada"}\n{"id":"002","name":"Ben"}\n'),
    ("customers.ndjson", b'{"id":"001","name":"Ada"}\n{"id":"002","name":"Ben"}'),
    ("customers.xlsx", workbook_bytes([["id", "name"], ["001", "Ada"], ["002", "Ben"]])),
])
def test_supported_formats_produce_same_preserved_table(filename, data):
    parsed = parse_upload(data, filename)
    assert table(parsed) == [["id", "name"], ["001", "Ada"], ["002", "Ben"]]
    assert parsed.manifest["version"] == 1
    assert parsed.manifest["row_count"] == 2
    assert parsed.manifest["column_count"] == 2
    assert parsed.preview.values.tolist() == [["001", "Ada"], ["002", "Ben"]]
    # The entire manifest is JSON-compatible for durable approval hashing.
    assert json.loads(json.dumps(parsed.manifest)) == parsed.manifest


def test_csv_retains_literals_blank_rows_and_multiline_source_positions():
    data = b'\xef\xbb\xbfid,note\r\n001,"hello\r\nworld"\r\n\r\n002,NA\r\n003,null\r\n'
    parsed = parse_upload(data, "data.CSV")
    assert table(parsed) == [["id", "note"], ["001", "hello\r\nworld"], ["", ""], ["002", "NA"], ["003", "null"]]
    assert parsed.manifest["source_locations"]["rows"] == [[2, 3], [4, 4], [5, 5], [6, 6]]
    assert "Retained 1" in parsed.manifest["warnings"][0]


def test_tsv_keeps_quoted_tabs_and_is_not_guessed_as_csv():
    parsed = parse_upload(b'id\tnote\n001\t"a\tb,c"\n', "source.tsv")
    assert table(parsed)[1] == ["001", "a\tb,c"]
    assert parsed.manifest["options"]["delimiter"] == "\t"


@pytest.mark.parametrize("data", [
    b"", b"id\n", b"id,id\na,b\n", b"id, \na,b\n", b"id,name\na\n",
    b"id,name\na,b,c\n", b'id,name\n"unfinished,b\n', b"id\n\xff\n", b"id\na\x00\n",
])
def test_malformed_delimited_input_is_rejected(data):
    with pytest.raises(ImportValidationError):
        parse_upload(data, "data.csv")


def test_json_scalar_types_missing_null_empty_and_exact_numeric_tokens():
    source = b'[{"id":"001","n":123456789012345678901234567890,"d":1.230000000000000000000001,"b":true,"v":null,"z":-0},{"id":"002","n":1E+100,"d":0.00000100,"b":false,"v":""},{"id":"003"}]'
    parsed = parse_upload(source, "data.json")
    assert table(parsed)[1] == ["001", "123456789012345678901234567890", "1.230000000000000000000001", "true", "", "-0"]
    assert table(parsed)[2][1:4] == ["1E+100", "0.00000100", "false"]
    assert parsed.manifest["numeric_lexemes"]["n"] == {"0": "123456789012345678901234567890", "1": "1E+100"}
    assert parsed.manifest["types"]["default"] == "string"
    types = parsed.manifest["types"]["by_column"]
    assert types["v"] == {"null": [0], "missing": [2]}
    assert types["b"] == {"boolean": [0, 1], "missing": [2]}
    assert parsed.manifest["source_locations"]["rows"] == [0, 1, 2]
    assert any("missing" in message and "null" in message for message in parsed.manifest["warnings"])


def test_json_empty_objects_remain_records_and_union_columns_are_stable():
    parsed = parse_upload(b'[{"b":"002"},{},{"a":"001","b":"003"}]', "data.json")
    assert table(parsed) == [["b", "a"], ["002", ""], ["", ""], ["003", "001"]]
    assert parsed.manifest["types"]["by_column"]["a"]["missing"] == [0, 1]


@pytest.mark.parametrize("source", [
    b'{"id":"001"}', b"[]", b"[{}]", b'[1]', b'[{"id":{}}]', b'[{"id":[]}]',
    b'[{"id":"001","id":"002"}]', b'[{"id":NaN}]', b'[{"id":Infinity}]', b'[{"id":-Infinity}]',
    b'[{" ":"x"}]', b'[{"id":"\\ud800"}]', b'[{"id":"\\u0000"}]', b'[{"id":01}]', b'[{"id":1',
])
def test_invalid_or_nested_json_is_rejected(source):
    with pytest.raises(ImportValidationError):
        parse_upload(source, "data.json")


@pytest.mark.parametrize("source", [b'\n{"a":1}', b'{"a":1}\n\n', b'{"a":1}\n \n{"a":2}', b'{"a":1}\n[]', b'{"a":1,"a":2}'])
def test_jsonl_never_silently_skips_blank_or_invalid_records(source):
    with pytest.raises(ImportValidationError):
        parse_upload(source, "data.jsonl")


def test_jsonl_keeps_unicode_line_separator_inside_string_and_line_provenance():
    parsed = parse_upload('{"a":"left\u2028right"}\r\n{"a":"last"}\r\n'.encode(), "data.ndjson")
    assert parsed.preview.iloc[0, 0] == "left\u2028right"
    assert parsed.manifest["source_locations"]["rows"] == [1, 2]


@pytest.mark.parametrize("filename", ["data.xls", "data.xlsm", "data.zip", "data.txt", "data.parquet", "data", ""])
def test_unsupported_extensions_are_rejected(filename):
    with pytest.raises(ImportValidationError):
        parse_upload(b"id\n001\n", filename)


def test_display_filename_cannot_become_storage_path():
    parsed = parse_upload(b"id\n001\n", r"C:\untrusted\path\data.CSV")
    assert parsed.manifest["original_filename"] == "data.CSV"
    with pytest.raises(ImportValidationError):
        parse_upload(b"id\n001\n", "bad\nname.csv")
    with pytest.raises(ImportValidationError):
        parse_upload(b"id\n001\n", "bad\ud800.csv")


def test_preview_is_bounded_without_truncating_import():
    source = ("id\n" + "\n".join(f"{i:04}" for i in range(125))).encode()
    parsed = parse_upload(source, "data.csv")
    assert len(parsed.preview) == 100
    assert parsed.manifest["row_count"] == 125
    assert len(table(parsed)) == 126


@pytest.mark.parametrize("setting,value,source,filename", [
    ("MAX_BYTES", 5, b"id\n001\n", "data.csv"),
    ("MAX_COLUMNS", 1, b"a,b\n1,2\n", "data.csv"),
    ("MAX_ROWS", 1, b"a\n1\n2\n", "data.csv"),
    ("MAX_CELLS", 3, b"a,b\n1,2\n3,4\n", "data.csv"),
    ("MAX_DECODED_BYTES", 5, b"a\n1234\n", "data.csv"),
    ("MAX_ROWS", 1, b'[{"a":1},{"a":2}]', "data.json"),
    ("MAX_COLUMNS", 1, b'[{"a":1},{"b":2}]', "data.json"),
    ("MAX_CELLS", 3, b'[{"a":1,"b":2},{"a":3}]', "data.json"),
    ("MAX_ROWS", 1, b'{"a":1}\n{"a":2}', "data.jsonl"),
    ("MAX_PARSE_SECONDS", -1, b"a\n1\n", "data.csv"),
])
def test_table_and_resource_limits(monkeypatch, setting, value, source, filename):
    monkeypatch.setattr(adapters, setting, value)
    with pytest.raises(ImportValidationError):
        parse_upload(source, filename)


def test_normalized_working_table_limit_applies_after_csv_escaping(monkeypatch):
    monkeypatch.setattr(adapters, "MAX_BYTES", 6)
    with pytest.raises(ImportValidationError, match="normalized table"):
        parse_upload(b"a\n\n\n", "data.csv")


def test_multiple_worksheets_require_explicit_selection_and_list_hidden_sheets():
    def configure(workbook):
        other = workbook.create_sheet("Hidden")
        other.append(["id", "name"])
        other.append(["002", "Ben"])
        other.sheet_state = "hidden"
    data = workbook_bytes(customize=configure)
    inspection = inspect_upload(data, "data.xlsx")
    assert inspection["sheets"] == ["Customers", "Hidden"]
    assert inspection["requires_sheet"] is True
    assert inspection["sheet_states"]["Hidden"] == "hidden"
    with pytest.raises(ImportValidationError, match="Select the worksheet"):
        parse_upload(data, "data.xlsx")
    parsed = parse_upload(data, "data.xlsx", sheet="Hidden")
    assert parsed.preview.values.tolist() == [["002", "Ben"]]
    assert parsed.manifest["sheet"] == "Hidden"
    assert any("hidden" in warning for warning in parsed.manifest["warnings"])
    with pytest.raises(ImportValidationError, match="does not exist"):
        parse_upload(data, "data.xlsx", sheet="Other")
    with pytest.raises(ImportValidationError, match="only to XLSX"):
        parse_upload(b"id\n001", "data.csv", sheet="Customers")


def test_xlsx_types_dates_and_text_identifiers_are_preserved():
    data = workbook_bytes([["id", "number", "active", "day", "stamp", "time"],
                           ["00001", 1.25, True, date(2026, 10, 25), datetime(2026, 10, 25, 11, 22, 33), time(12, 30)]])
    parsed = parse_upload(data, "data.xlsx")
    assert table(parsed)[1] == ["00001", "1.25", "true", "2026-10-25", "2026-10-25T11:22:33", "12:30:00"]
    types = parsed.manifest["types"]["by_column"]
    assert "id" not in types
    assert types["number"] == {"number": [0]}
    assert types["day"] == {"date": [0]}
    assert parsed.manifest["source_locations"] == {"kind": "worksheet", "sheet": "Customers", "rows": [2], "columns": ["A", "B", "C", "D", "E", "F"]}


def test_xlsx_numeric_tokens_do_not_round_through_openpyxl_float():
    data = workbook_bytes([["id", "amount"], ["001", 1.25]])
    data = rewrite_zip(data, {"xl/worksheets/sheet1.xml": lambda xml: xml.replace(b"<v>1.25</v>", b"<v>1.234567890123456789</v>")})
    parsed = parse_upload(data, "data.xlsx")
    assert table(parsed)[1][1] == "1.234567890123456789"
    assert parsed.manifest["numeric_lexemes"]["amount"]["0"] == "1.234567890123456789"


@pytest.mark.parametrize("format_code", ["00000", '"ID "00000', "000-000", "00000;[Red]-00000"])
def test_xlsx_zero_padded_identifiers_require_text_values(format_code):
    def configure(workbook):
        workbook.active["A2"] = 123
        workbook.active["A2"].number_format = format_code
    with pytest.raises(ImportValidationError, match="zero-padded"):
        parse_upload(workbook_bytes(customize=configure), "data.xlsx")


def test_xlsx_numeric_display_format_does_not_change_stored_value():
    def configure(workbook):
        workbook.active["A2"] = 0.125
        workbook.active["A2"].number_format = "0.00%"
    parsed = parse_upload(workbook_bytes(customize=configure), "data.xlsx")
    assert parsed.preview.iloc[0, 0] == "0.125"
    assert any("display formatting" in warning for warning in parsed.manifest["warnings"])


def test_xlsx_blank_and_hidden_rows_are_retained():
    def configure(workbook):
        worksheet = workbook.active
        worksheet["A4"] = "003"
        worksheet["B4"] = "Cara"
        worksheet.row_dimensions[2].hidden = True
        worksheet.column_dimensions["B"].hidden = True
    parsed = parse_upload(workbook_bytes(customize=configure), "data.xlsx")
    assert table(parsed)[1:] == [["001", "Ada"], ["", ""], ["003", "Cara"]]
    assert parsed.manifest["source_locations"]["rows"] == [2, 3, 4]
    assert parsed.manifest["types"]["by_column"]["id"]["blank"] == [1]
    assert any("Hidden rows/columns" in warning for warning in parsed.manifest["warnings"])
    assert any("blank worksheet row" in warning for warning in parsed.manifest["warnings"])


@pytest.mark.parametrize("configure,match", [
    (lambda w: setattr(w.active["A2"], "value", "=1+1"), "formulas"),
    (lambda w: w.active.merge_cells("A2:B2"), "merged cells"),
    (lambda w: setattr(w.active["A2"], "value", "#DIV/0!"), "Excel error"),
    (lambda w: setattr(w.active["A1"], "value", 123), "text column headers"),
    (lambda w: setattr(w.active["B1"], "value", "id"), "unique"),
    (lambda w: setattr(w.active["A2"], "hyperlink", "https://example.com"), "external links"),
])
def test_xlsx_ambiguous_or_active_content_is_rejected(configure, match):
    with pytest.raises(ImportValidationError, match=match):
        parse_upload(workbook_bytes(customize=configure), "data.xlsx")


def test_xlsx_cached_formula_is_not_treated_as_a_value():
    data = workbook_bytes()
    data = rewrite_zip(data, {"xl/worksheets/sheet1.xml": lambda xml: xml.replace(b'<row r="2">', b'<row r="2"><c r="C2"><f>1+1</f><v>2</v></c>')})
    with pytest.raises(ImportValidationError, match="formulas"):
        inspect_upload(data, "data.xlsx")


@pytest.mark.parametrize("extra", ["xl/vbaProject.bin", "xl/externalLinks/externalLink1.xml", "xl/connections.xml", "xl/activeX/activeX1.xml", "xl/embeddings/oleObject1.bin", "../escape.xml"])
def test_xlsx_zip_rejects_macros_connections_and_traversal(extra):
    with pytest.raises(ImportValidationError):
        inspect_upload(rewrite_zip(workbook_bytes(), extras={extra: b""}), "data.xlsx")


def test_xml_entities_are_rejected_before_openpyxl():
    xml = b'<!DOCTYPE Types [<!ENTITY x "expanded">]><Types>&x;</Types>'
    data = rewrite_zip(workbook_bytes(), {"[Content_Types].xml": lambda _: xml})
    with pytest.raises(ImportValidationError, match="unsafe entity/DTD"):
        inspect_upload(data, "data.xlsx")


@pytest.mark.parametrize("setting,value,match", [
    ("MAX_ZIP_MEMBERS", 1, "too many ZIP"),
    ("MAX_DECODED_BYTES", 100, "ZIP expansion"),
    ("MAX_XML_DEPTH", 2, "deeply nested"),
    ("MAX_COLUMNS", 1, "200 columns"),
    ("MAX_CELLS", 1, "data cells"),
])
def test_xlsx_resource_bounds(monkeypatch, setting, value, match):
    data = workbook_bytes()
    monkeypatch.setattr(adapters, setting, value)
    with pytest.raises(ImportValidationError, match=match):
        parse_upload(data, "data.xlsx")


def test_xlsx_ignores_false_dimension_hint_but_rejects_out_of_bounds_real_cells():
    data = workbook_bytes()
    false_dimension = rewrite_zip(data, {"xl/worksheets/sheet1.xml": lambda xml: xml.replace(b'ref="A1:B2"', b'ref="A1:XFD1048576"')})
    assert parse_upload(false_dimension, "data.xlsx").manifest["row_count"] == 1
    out_of_bounds = rewrite_zip(data, {"xl/worksheets/sheet1.xml": lambda xml: xml.replace(b'r="A2"', b'r="XFD1048576"')})
    with pytest.raises(ImportValidationError):
        parse_upload(out_of_bounds, "data.xlsx")


def test_xlsx_rejects_duplicate_zip_members():
    with pytest.warns(UserWarning, match="Duplicate name"):
        data = rewrite_zip(workbook_bytes(), extras={"xl/workbook.xml": b"<workbook/>"})
    with pytest.raises(ImportValidationError, match="duplicate or unsafe"):
        inspect_upload(data, "data.xlsx")


def test_xlsx_rejects_mismatched_cell_and_row_positions():
    data = rewrite_zip(workbook_bytes(), {"xl/worksheets/sheet1.xml": lambda xml: xml.replace(b'r="A2"', b'r="A3"')})
    with pytest.raises(ImportValidationError, match="must match their row"):
        parse_upload(data, "data.xlsx")


@pytest.mark.parametrize("date_format", ["yyyy-mm-dd", "YYYY-MM-DD"])
def test_xlsx_nonexistent_1900_leap_day_is_not_silently_corrected(date_format):
    def configure(workbook):
        workbook.active["A2"] = 60
        workbook.active["A2"].number_format = date_format
    with pytest.raises(ImportValidationError, match="nonexistent 1900-02-29"):
        parse_upload(workbook_bytes(customize=configure), "data.xlsx")


def test_native_excel_iso_datetime_preserves_timezone():
    def configure(workbook):
        workbook.iso_dates = True
    data = workbook_bytes([["id", "stamp"], ["001", datetime(2026, 10, 25, 11, 22, 33)]], configure)
    data = rewrite_zip(data, {"xl/worksheets/sheet1.xml": lambda xml: xml.replace(b"2026-10-25T11:22:33", b"2026-10-25T11:22:33+05:30")})
    assert parse_upload(data, "data.xlsx").preview.iloc[0, 1] == "2026-10-25T11:22:33+05:30"


def test_native_excel_date_with_invalid_trailing_text_is_rejected():
    def configure(workbook):
        workbook.iso_dates = True
    data = workbook_bytes([["id", "day"], ["001", date(2026, 10, 25)]], configure)
    data = rewrite_zip(data, {"xl/worksheets/sheet1.xml": lambda xml: xml.replace(b"2026-10-25", b"2026-10-25invalid")})
    with pytest.raises(ImportValidationError, match="invalid ISO"):
        parse_upload(data, "data.xlsx")


def test_excel_date_only_display_does_not_discard_stored_time():
    def configure(workbook):
        workbook.active["B2"].number_format = "yyyy-mm-dd"
    data = workbook_bytes([["id", "stamp"], ["001", datetime(2026, 10, 25, 11, 22, 33)]], configure)
    parsed = parse_upload(data, "data.xlsx")
    assert parsed.preview.iloc[0, 1] == "2026-10-25T11:22:33"
    assert parsed.manifest["types"]["by_column"]["stamp"] == {"datetime": [0]}


@pytest.mark.parametrize("data", [b"not a zip", b"PK\x03\x04bad", b"id,name\n001,Ada\n"])
def test_mismatched_or_corrupt_xlsx_is_rejected(data):
    with pytest.raises(ImportValidationError):
        parse_upload(data, "data.xlsx")
